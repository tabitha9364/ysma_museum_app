from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    r"C:\Users\ISMS\Downloads\ISMS-Project-Writing-Format.docx"
)
OUT = ROOT / "docs" / "YSMA_Final_Year_Project_Chapters_4_5_updated.docx"


def clear_document(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def apply_run_font(run, size=12, bold=False, italic=False, color=None, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_styles(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    paragraph_format = normal.paragraph_format
    paragraph_format.first_line_indent = Inches(0.5)
    paragraph_format.line_spacing = 2
    paragraph_format.space_after = Pt(0)
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.bold = True
    styles["Heading 3"].font.italic = True
    styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_para(document: Document, text: str = "", *, style=None, first_line=True, justify=True):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Inches(0.5 if first_line else 0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    if text:
        run = paragraph.add_run(text)
        apply_run_font(run)
    return paragraph


def add_heading(document: Document, text: str, level: int):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    apply_run_font(run, bold=True, italic=(level == 3))
    return paragraph


def add_caption(document: Document, label: str, title: str):
    p1 = document.add_paragraph()
    p1.paragraph_format.first_line_indent = Inches(0)
    p1.paragraph_format.line_spacing = 1
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(label)
    apply_run_font(r1, bold=True)

    p2 = document.add_paragraph()
    p2.paragraph_format.first_line_indent = Inches(0)
    p2.paragraph_format.line_spacing = 1
    p2.paragraph_format.space_after = Pt(6)
    r2 = p2.add_run(title)
    apply_run_font(r2, italic=True)


def add_note(document: Document, text: str):
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    apply_run_font(r, size=10)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1
        run = p.add_run(header)
        apply_run_font(run, size=10, bold=True)
        set_cell_shading(hdr[i], "F2F2F2")
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            hdr[i].width = Inches(widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            apply_run_font(run, size=10)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Inches(widths[i])

    document.add_paragraph()
    return table


def add_bullets(document: Document, items: list[str]):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.line_spacing = 2
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(item)
        apply_run_font(r)


def add_numbered(document: Document, items: list[str]):
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.line_spacing = 2
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(item)
        apply_run_font(r)


def add_code(document: Document, code: str):
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(8)
    for line in code.strip("\n").splitlines():
        r = p.add_run(line.rstrip())
        apply_run_font(r, size=9, font="Courier New")
        p.add_run().add_break()


def add_reference(document: Document, text: str):
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    apply_run_font(r)


def build_document() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(str(TEMPLATE)) if TEMPLATE.exists() else Document()
    clear_document(document)
    setup_styles(document)

    add_heading(document, "CHAPTER FOUR", 1)
    add_heading(document, "IMPLEMENTATION", 1)
    add_para(
        document,
        "This chapter presents the implementation of YSMA Smart Experience, a mobile software production project designed to enhance visitor engagement at the Yemisi Shyllon Museum of Art. In line with the production-type project guidelines of the Design and New Media Department, the chapter explains the system requirements, execution process and major outputs of the implemented application. The implementation covers the mobile interface, artwork recognition model, augmented reality style feedback, audio guide, museum navigation, accessibility settings, user activity tracking and management analytics dashboard.",
    )

    add_heading(document, "4.1 System Requirements", 2)
    add_para(
        document,
        "The system requirements define the resources needed to develop, install, run and maintain the application. Since the project is a software application, the requirements are grouped into hardware requirements, software requirements and interface requirements. These requirements cover both the development environment used by the project developer and the operational environment required by museum visitors and administrators.",
    )

    add_heading(document, "4.1.1 Hardware Requirements", 3)
    add_para(
        document,
        "The hardware requirements represent the physical devices needed for successful implementation and use of the application. The application is intended to run on an Android smartphone because the main visitor functions depend on camera input, mobile display output and audio playback.",
    )
    add_caption(document, "Table 4.1", "Hardware requirements for YSMA Smart Experience")
    add_table(
        document,
        ["Hardware item", "Minimum requirement", "Recommended requirement", "Purpose"],
        [
            [
                "Development computer",
                "Intel Core i3 or equivalent processor, 8 GB RAM and 20 GB free storage",
                "Intel Core i5/i7 or equivalent processor, 16 GB RAM and SSD storage",
                "Used for coding, model integration, application building, testing and debugging.",
            ],
            [
                "Android mobile device",
                "Android 8.0 device with 3 GB RAM and a functional rear camera",
                "Android 10 or later device with 4 GB RAM or more and a high-quality autofocus rear camera",
                "Used by visitors to install the application, scan artworks, read information and navigate the museum.",
            ],
            [
                "Camera",
                "Rear camera with basic autofocus",
                "Rear camera with good focus, exposure control and low-light performance",
                "Captures artwork images for the TensorFlow Lite recognition model.",
            ],
            [
                "Storage",
                "At least 500 MB free mobile storage",
                "1 GB or more free mobile storage",
                "Stores the application, cached artwork images, model files, preferences and local activity records.",
            ],
            [
                "Network access",
                "Mobile data or Wi-Fi connection",
                "Stable 4G/5G or museum Wi-Fi connection",
                "Fetches Supabase artwork records, images, authentication data and optional cloud analytics.",
            ],
            [
                "Audio output",
                "Phone speaker",
                "Phone speaker or headphones",
                "Plays artwork audio guides and voice navigation instructions.",
            ],
        ],
        widths=[1.35, 1.75, 1.75, 1.65],
    )
    add_note(
        document,
        "Note. The mobile camera and audio output are essential because artwork recognition, audio narration and voice navigation are central functions of the application.",
    )

    add_heading(document, "4.1.2 Software Requirements", 3)
    add_para(
        document,
        "The software requirements include the operating systems, frameworks, programming tools, runtime libraries and cloud services used to implement the system. Flutter was selected because it supports rapid mobile interface development with a single Dart codebase, while Supabase provides backend storage and authentication. TensorFlow Lite was used to deploy the image classification model on the mobile device.",
    )
    add_caption(document, "Table 4.2", "Software requirements for system development and execution")
    add_table(
        document,
        ["Software item", "Description", "Use in the project"],
        [
            [
                "Flutter SDK",
                "Cross-platform mobile development framework",
                "Used to build the Android application screens, navigation, state management and interactive user interface.",
            ],
            [
                "Dart",
                "Programming language for Flutter",
                "Used to implement application logic, services, models, screen widgets and asynchronous operations.",
            ],
            [
                "Android Studio or Visual Studio Code",
                "Integrated development environment",
                "Used for coding, debugging, emulator testing and application build management.",
            ],
            [
                "Android SDK",
                "Mobile build and deployment toolkit",
                "Used to compile and package the Flutter application for Android devices.",
            ],
            [
                "Supabase",
                "Backend-as-a-service platform",
                "Stores artwork records and supports user authentication and optional central analytics storage.",
            ],
            [
                "TensorFlow Lite",
                "Mobile machine learning runtime",
                "Runs the trained artwork recognition model directly on the visitor's mobile device.",
            ],
            [
                "Camera package",
                "Flutter camera integration package",
                "Captures camera images for scanning and recognition.",
            ],
            [
                "just_audio package",
                "Audio playback package",
                "Plays artwork audio guide files in the artwork detail module.",
            ],
            [
                "shared_preferences package",
                "Local key-value storage package",
                "Stores user preferences, favourites, recent activity and local analytics fallback events.",
            ],
            [
                "cached_network_image package",
                "Image loading and caching package",
                "Displays artwork images efficiently and reduces repeated network loading.",
            ],
        ],
        widths=[1.7, 2.05, 2.75],
    )
    add_note(
        document,
        "Note. The system uses a hybrid data approach: museum content is fetched from Supabase while selected preferences and activity records are cached locally for responsiveness.",
    )

    add_heading(document, "4.1.3 Interface Requirements", 3)
    add_para(
        document,
        "The interface requirements describe the interaction points through which visitors, administrators and device services communicate with the system. The application must provide a clear visitor interface, reliable camera input, readable content screens, audio output and a management view for activity analytics.",
    )
    add_caption(document, "Table 4.3", "Interface requirements and expected behaviour")
    add_table(
        document,
        ["Interface", "Requirement", "Expected output"],
        [
            [
                "Visitor authentication interface",
                "The application must allow a visitor to sign in and preserve a recognizable visitor identity.",
                "The visitor's name is displayed on the home and profile screens after authentication.",
            ],
            [
                "Home interface",
                "The application must present search, recommended artworks and quick access actions.",
                "Visitors can open scanning, the museum map, artwork spotlight and artwork previews.",
            ],
            [
                "Camera scanning interface",
                "The application must access the rear camera and send captured frames to the recognition service.",
                "Only confident matches from the 50 artwork classes proceed to the success output; weak or unrelated inputs are rejected.",
            ],
            [
                "Artwork information interface",
                "The application must display image, title, artist, year, description, audio and related artworks.",
                "Visitors receive readable artwork interpretation and can move between description, audio and related tabs.",
            ],
            [
                "Navigation interface",
                "The application must present gallery zones, route paths, destination pins and voice guidance.",
                "Visitors can locate a selected artwork and hear spoken route instructions.",
            ],
            [
                "Accessibility interface",
                "The application must support visitor comfort preferences such as text size, theme and audio settings.",
                "The user interface becomes more inclusive for visitors with different reading and display needs.",
            ],
            [
                "Admin analytics interface",
                "The application must summarize user activity for museum management.",
                "Management can view total artwork views, shares, favourites, visitor count, top artworks and recent activity.",
            ],
        ],
        widths=[1.55, 2.45, 2.5],
    )
    add_note(
        document,
        "Note. Interface requirements are included because the system depends not only on software execution but also on clear human interaction, device input and management reporting.",
    )

    add_heading(document, "4.2 System Execution/Output", 2)
    add_para(
        document,
        "System execution refers to the process through which the developed mobile application is launched, operated and used to produce meaningful outputs. The system is operationalized as an Android mobile application connected to Supabase and supported by an on-device TensorFlow Lite model. Once installed, the visitor launches the application, signs in, opens the home screen, scans artworks, views AR-style confirmation feedback, reads or listens to interpretation, locates artworks in the museum and generates activity records for analytics.",
    )

    add_heading(document, "4.2.1 Application Start-up and Operational Flow", 3)
    add_para(
        document,
        "To execute the system, the mobile device is switched on, connected to the internet and the YSMA Smart Experience application is opened from the device launcher. During start-up, the application initializes Supabase and loads the TensorFlow Lite artwork recognition model. The splash screen then introduces the application and provides a tap-to-continue action. After authentication, the system retrieves artwork data and displays the home screen.",
    )
    add_numbered(
        document,
        [
            "The visitor switches on the Android device and connects to mobile data or Wi-Fi.",
            "The visitor opens YSMA Smart Experience from the mobile device.",
            "The application initializes Supabase, loads the artwork model and displays the splash screen.",
            "The visitor taps the continue action and signs in.",
            "The system fetches artwork records and presents the home screen.",
            "The visitor selects a task such as scanning an artwork, searching the collection, opening the map or viewing the spotlight.",
            "The selected module processes the request and produces the relevant output on screen.",
            "Visitor activity such as viewing, sharing or favouriting an artwork is recorded for profile and analytics output.",
        ],
    )

    add_heading(document, "4.2.2 Authentication and Home Screen Output", 3)
    add_para(
        document,
        "The authentication module allows the visitor to access the application with a personalized identity. When a visitor signs in, the application reads the user's display name or email metadata and uses it to personalize the home and profile screens. This output supports a more user-centred museum experience because the visitor is not treated as an anonymous user after sign-in.",
    )
    add_para(
        document,
        "The home screen is the main control point of the application. It displays the visitor greeting, search bar, recommended artworks and quick action cards. The search function filters artworks by title, artist, year, tag and location. The output of this module is a navigable digital collection interface through which visitors can move quickly to scanning, map navigation, artwork spotlight or individual artwork pages.",
    )

    add_heading(document, "4.2.3 Artwork Recognition and AR Confirmation Output", 3)
    add_para(
        document,
        "The artwork recognition module uses the mobile camera and the TensorFlow Lite model to classify captured images. The model output is not accepted automatically. The application checks that the prediction is not the unknown class, that the model confidence is high, that the difference between the first and second predictions is wide enough, that multiple image crops agree with the same result and that the predicted label matches an artwork title in the application data. In the implemented system, the confidence threshold is 0.90, the confidence gap threshold is 0.24 and the crop agreement threshold is 0.80. These checks are important because the system should not force a random object into one of the museum artwork classes.",
    )
    add_para(
        document,
        "If the captured image fails the recognition checks, the output is an 'Artwork not recognized' status and the scan continues. If the image passes the checks, the system displays a recognition success overlay. This overlay contains an animated confirmation card, glowing recognition effects, haptic feedback, artwork image preview, progress indicator and a compact approval mascot that gives a thumbs-up. The mascot appears at the same time as the recognition animation to communicate that the artwork has been accepted without overloading the screen.",
    )

    add_heading(document, "4.2.4 Artwork Preview, Details and Audio Output", 3)
    add_para(
        document,
        "After a successful match, the application opens the artwork preview screen. This screen presents the recognized artwork image, title, artist, year and museum location. It also provides the major visitor actions: Play Audio, Read Info and Locate in Museum. These outputs allow the visitor to move from recognition to interpretation, listening or navigation without returning to the home screen.",
    )
    add_para(
        document,
        "The artwork detail screen contains the description, audio and related artwork tabs. The description tab presents interpretive text, the audio tab provides playback controls for the audio guide and the related tab recommends artworks connected by artist, year or tag. This structure supports both quick reading and deeper exploration.",
    )

    add_heading(document, "4.2.5 Museum Map and Voice Navigation Output", 3)
    add_para(
        document,
        "The museum map module operationalizes virtual navigation within the application. It divides the museum experience into named zones such as Heritage Hall, Masters Gallery, Sculpture Court, Performance Wing and Contemporary Wing. When a visitor selects a gallery or an artwork, the system displays a highlighted route from the entrance to the destination. The selected artwork is represented with a pulsing pin and a list of route steps.",
    )
    add_para(
        document,
        "Voice navigation is executed through the voice navigation service. The service first attempts to call the Android native text-to-speech channel. If the native voice channel is unavailable, it falls back to Flutter's accessibility announcement service, allowing the route instruction to be announced through device accessibility audio. The result is a navigation output similar in purpose to standard navigation systems: a selected destination, a visible route and spoken route guidance.",
    )

    add_heading(document, "4.2.6 Artwork Spotlight Output", 3)
    add_para(
        document,
        "The artwork spotlight module highlights a selected artwork and presents a clearly labelled Fun Facts section. Unlike the normal description field, the spotlight facts are written from public online art references such as Google Arts and Culture. This makes the spotlight output more engaging because it gives visitors extra contextual facts instead of repeating only the artwork information already stored in the application database.",
    )

    add_heading(document, "4.2.7 Profile and Admin Analytics Output", 3)
    add_para(
        document,
        "The profile screen records user-facing activity such as viewed artworks, shares and favourites. These activities are stored locally so that the visitor can see a recent activity history and daily journey summary. The same activity events are also used by the admin analytics module.",
    )
    add_para(
        document,
        "The admin analytics dashboard is designed for museum management. It displays total artwork views, shares, favourites, visitor count, top viewed artworks and recent visitor activity. The system stores analytics immediately on the device and can also send events to a Supabase table named admin_activity_events when that table is created. This gives the museum a simple foundation for data-informed decisions about visitor engagement and artwork popularity.",
    )

    add_heading(document, "4.2.8 System Testing and Observed Outputs", 3)
    add_para(
        document,
        "The implemented system was tested by opening the main modules and checking whether each module produced the expected output. The tests focused on the visitor journey, recognition process, navigation support, audio output and analytics recording.",
    )
    add_caption(document, "Table 4.4", "System test summary")
    add_table(
        document,
        ["Test area", "Expected output", "Observed implementation output"],
        [
            [
                "Splash screen",
                "The application opens and allows the visitor to continue.",
                "The splash screen displays the application identity and opens the authentication flow.",
            ],
            [
                "Authentication",
                "The visitor can sign in and receive a personalized experience.",
                "The visitor name is shown on the home and profile screens after sign-in.",
            ],
            [
                "Artwork scanning",
                "The system recognizes only trained artworks and rejects unrelated inputs.",
                "The scan module applies confidence, confidence gap, crop agreement and title-matching checks before accepting a result.",
            ],
            [
                "Recognition confirmation",
                "A clear success output appears after recognition.",
                "The system displays an animated AR-style card, progress feedback and approval mascot before opening the artwork preview.",
            ],
            [
                "Audio guide",
                "The visitor can play and pause artwork narration.",
                "The artwork detail screen provides audio playback controls using the audio package.",
            ],
            [
                "Map navigation",
                "The visitor can locate an artwork in the museum.",
                "The map screen displays zones, route line, route steps, destination pin and voice navigation.",
            ],
            [
                "Admin analytics",
                "Management can view visitor activity summaries.",
                "The dashboard displays views, shares, favourites, visitor count, top artworks and recent events.",
            ],
        ],
        widths=[1.6, 2.25, 2.65],
    )

    document.add_page_break()
    add_heading(document, "CHAPTER FIVE", 1)
    add_heading(document, "SUMMARY, CONCLUSION AND RECOMMENDATIONS", 1)

    add_heading(document, "5.1 Summary", 2)
    add_para(
        document,
        "This project focused on the design and implementation of YSMA Smart Experience, a smart museum mobile application for the Yemisi Shyllon Museum of Art. The project was developed as a production-type software project intended to solve practical visitor experience problems such as limited interactivity, lack of digital artwork recognition, weak navigation support, limited accessibility options and absence of basic visitor analytics.",
    )
    add_para(
        document,
        "The implemented system combines mobile application development, image recognition, AR-style visual feedback, audio guidance, museum navigation, accessibility support and analytics. Flutter and Dart were used for the mobile application, Supabase was used for authentication and artwork data storage, TensorFlow Lite was used for on-device image recognition and local storage was used for user preferences and activity records.",
    )
    add_para(
        document,
        "The application enables visitors to sign in, browse artworks, search the collection, scan artworks with the phone camera, receive animated recognition feedback, read artwork descriptions, play audio guides, locate artworks on the museum map, view fun facts in the spotlight section and track recent activity. For museum management, the admin analytics dashboard provides a simple reporting layer that summarizes engagement through views, shares, favourites, visitor count, top artworks and recent events.",
    )
    add_para(
        document,
        "Overall, the system demonstrates how a mobile application can support museum interpretation by connecting physical artworks with digital content and visitor-centred interaction. The implementation also shows that technology can improve the museum experience without removing the value of seeing artworks physically in the gallery.",
    )

    add_heading(document, "5.2 Conclusion", 2)
    add_para(
        document,
        "The project successfully achieved its aim of developing a mobile application that enhances museum experience through artwork recognition, digital interpretation, navigation and analytics. The completed application provides a practical visitor guide that can identify selected artworks, present relevant information, support audio learning and guide visitors to artwork locations within the museum.",
    )
    add_para(
        document,
        "A major contribution of the implementation is the stricter recognition workflow. Instead of accepting every camera input as one of the 50 artwork classes, the system applies confidence, confidence gap, crop agreement, unknown-class and database-title checks before accepting a prediction. This makes the recognition process more reliable for a museum setting because unrelated objects are rejected rather than falsely matched.",
    )
    add_para(
        document,
        "The navigation module also strengthens the usefulness of the application by turning the museum map into an interactive guide. The visitor can select a destination, view a highlighted route and listen to spoken directions. The use of Flutter's semantics announcement as a fallback also improves accessibility when the native voice channel is unavailable.",
    )
    add_para(
        document,
        "In conclusion, YSMA Smart Experience meets the production objective of implementing a digital solution to a defined museum communication and visitor engagement problem. The system is functional as a prototype and provides a foundation that can be extended into a full museum deployment with additional training data, stronger analytics infrastructure and real indoor positioning.",
    )

    add_heading(document, "5.3 Recommendations", 2)
    add_para(
        document,
        "Although the system meets its major objectives, the following recommendations are made for future improvement:",
    )
    add_numbered(
        document,
        [
            "The artwork recognition model should be retrained with more real smartphone images taken inside the museum under different lighting conditions, camera angles, distances and visitor movement conditions.",
            "The unknown class should be expanded with more negative examples, including random objects, people, walls, floors, labels and non-collection images, so that the system becomes better at rejecting unrelated scans.",
            "The museum map should be connected to a verified architectural floor plan and improved with indoor positioning technologies such as Bluetooth beacons, QR markers or Wi-Fi positioning.",
            "The admin analytics dashboard should be connected to a secured Supabase table with role-based access control so that only authorized museum staff can view management data.",
            "The audio guide feature should be expanded to include multiple languages and richer narration for visitors with different cultural and language backgrounds.",
            "The AR component can be extended beyond the current AR-style overlay to include true ARCore or ARKit placement, interactive labels and contextual annotations in the camera view.",
            "Offline caching should be strengthened so that visitors can continue browsing saved artworks, listening to cached audio and using basic map functions when internet connectivity is weak.",
            "Further usability testing should be conducted with students, museum visitors, curators and accessibility-sensitive users to refine the interface before public deployment.",
        ],
    )

    document.add_page_break()
    add_heading(document, "REFERENCES", 1)
    references = [
        "Design and New Media Department, School of Media and Communication, Pan-Atlantic University. (2021/2022). Final year project guidelines: B.Sc. Information Science and Media Studies (ISMS). Pan-Atlantic University.",
        "Flutter. (n.d.). Flutter documentation. https://docs.flutter.dev/",
        "Google. (n.d.). Teachable Machine. https://teachablemachine.withgoogle.com/",
        "Google Arts & Culture. (n.d.). Yemisi Shyllon Museum of Art. https://artsandculture.google.com/partner/yemisi-shyllon-museum-of-art",
        "Supabase. (n.d.). Supabase documentation. https://supabase.com/docs",
        "TensorFlow. (n.d.). TensorFlow Lite. https://www.tensorflow.org/lite",
    ]
    for reference in references:
        add_reference(document, reference)

    document.add_page_break()
    add_heading(document, "APPENDICES", 1)

    add_heading(document, "Appendix A: System Flowchart", 2)
    add_code(
        document,
        """
Start
  |
Launch YSMA Smart Experience
  |
Display Splash Screen
  |
Sign In / Sign Up
  |
Load Artwork Records and Recognition Model
  |
Display Home Screen
  |--------------------------|--------------------------|
Scan Artwork              Museum Map              Artwork Spotlight
  |                          |                          |
Capture Image             Select Destination       Display Fun Facts
  |
Run TensorFlow Lite Prediction
  |
Check confidence, gap, agreement, unknown class and artwork title
  |
Known artwork match? ---- No ----> Show "Artwork not recognized"
  |
 Yes
  |
Show AR Confirmation Overlay and Approval Mascot
  |
Open Artwork Preview
  |
Read Info / Play Audio / Locate in Museum / Favourite / Share
  |
Store Visitor Activity
  |
Profile Statistics and Admin Analytics Dashboard
  |
End
""",
    )

    add_heading(document, "Appendix B: Major Program Source Code Files", 2)
    add_table(
        document,
        ["Source file", "Main purpose"],
        [
            ["lib/main.dart", "Initializes Supabase, loads the recognition model and starts the Flutter application."],
            ["lib/services/artwork_detection_service.dart", "Loads the TensorFlow Lite model and performs artwork prediction with unknown rejection checks."],
            ["lib/screens/scan_screen.dart", "Handles camera scanning, automatic detection, recognition success overlay and preview navigation."],
            ["lib/screens/artwork_preview_screen.dart", "Displays recognized artwork summary and actions for audio, reading and museum location."],
            ["lib/screens/artwork_detail_screen.dart", "Displays description, audio guide and related artworks."],
            ["lib/screens/museum_map_screen.dart", "Provides gallery zones, route drawing, destination selection and voice navigation."],
            ["lib/services/voice_navigation_service.dart", "Calls Android text-to-speech and falls back to Flutter semantics announcement for route guidance."],
            ["lib/screens/admin_analytics_screen.dart", "Displays management analytics for views, shares, favourites, top artworks and recent activity."],
            ["lib/services/admin_analytics_service.dart", "Stores and retrieves visitor activity events for dashboard reporting."],
        ],
        widths=[2.45, 4.05],
    )

    add_heading(document, "Appendix C: Program Source Code Excerpts", 2)
    add_para(document, "Recognition acceptance logic:", first_line=False)
    add_code(
        document,
        """
if (result.isUnknown ||
    predictedKey == 'unknown' ||
    result.confidence < _minimumPredictionConfidence ||
    confidenceGap < _minimumPredictionGap ||
    result.agreement < _minimumPredictionAgreement) {
  _setScanStatus('Artwork not recognized. Keep scanning a museum artwork.');
  return;
}
""",
    )
    add_para(document, "Voice navigation fallback logic:", first_line=False)
    add_code(
        document,
        """
try {
  final spoken = await _channel.invokeMethod<bool>(
    'speak',
    {'text': text},
  );
  if (spoken == true) return true;
} catch (error) {
  debugPrint('Native voice navigation failed: $error');
}

await SemanticsService.sendAnnouncement(view, text, textDirection);
return false;
""",
    )
    add_para(document, "Admin analytics event storage:", first_line=False)
    add_code(
        document,
        """
static Future<void> logEvent({
  required String type,
  required String artworkTitle,
}) async {
  final user = Supabase.instance.client.auth.currentUser;
  final event = AdminActivityEvent(
    type: type,
    artworkTitle: artworkTitle,
    visitorName: _displayNameForUser(user),
    visitorEmail: user?.email ?? '',
    visitorId: user?.id ?? '',
    createdAt: DateTime.now(),
  );

  await _storeLocalEvent(event);
}
""",
    )

    add_heading(document, "Appendix D: Suggested Analytics Table", 2)
    add_code(
        document,
        """
create table admin_activity_events (
  id bigint generated by default as identity primary key,
  event_type text not null,
  artwork_title text not null,
  visitor_name text,
  visitor_email text,
  visitor_id uuid,
  created_at timestamptz not null default now()
);
""",
    )

    document.save(str(OUT))


if __name__ == "__main__":
    build_document()
    print(OUT)
