from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\ISMS\Downloads\chapter 3.docx")
OUT = ROOT / "docs" / "chapter 3 updated final.docx"
DIAGRAM_DIR = ROOT / "docs" / "chapter3_flowcharts"


PALETTE = {
    "ink": (21, 35, 52),
    "muted": (81, 96, 114),
    "line": (197, 211, 225),
    "gold": (245, 181, 35),
    "blue": (69, 182, 254),
    "green": (54, 179, 126),
    "red": (245, 92, 112),
    "purple": (139, 92, 246),
    "surface": (248, 250, 252),
    "white": (255, 255, 255),
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


TITLE_FONT = font(36, True)
BOX_FONT = font(24, True)
BODY_FONT = font(20)
SMALL_FONT = font(17)


def text_size(draw: ImageDraw.ImageDraw, text: str, chosen_font: ImageFont.ImageFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=chosen_font)
    return box[2] - box[0], box[3] - box[1]


def wrap_label(text: str, width_chars: int) -> list[str]:
    lines: list[str] = []
    for part in text.split("\n"):
        lines.extend(textwrap.wrap(part, width=width_chars) or [""])
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    lines: list[str],
    chosen_font: ImageFont.ImageFont,
    fill: tuple[int, int, int] = PALETTE["ink"],
    spacing: int = 5,
) -> None:
    x1, y1, x2, y2 = rect
    heights = [text_size(draw, line, chosen_font)[1] for line in lines]
    total_height = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + ((y2 - y1 - total_height) / 2)
    for line, height in zip(lines, heights):
        width, _ = text_size(draw, line, chosen_font)
        draw.text((x1 + ((x2 - x1 - width) / 2), y), line, font=chosen_font, fill=fill)
        y += height + spacing


def box(
    draw: ImageDraw.ImageDraw,
    rect: tuple[int, int, int, int],
    title: str,
    body: str = "",
    fill: tuple[int, int, int] = PALETTE["white"],
    outline: tuple[int, int, int] = PALETTE["line"],
    accent: tuple[int, int, int] = PALETTE["gold"],
    wrap: int = 26,
) -> None:
    draw.rounded_rectangle(rect, radius=22, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle((x1, y1, x1 + 14, y2), radius=10, fill=accent)
    title_lines = wrap_label(title, wrap)
    if body:
        body_lines = wrap_label(body, wrap + 8)
        title_height = len(title_lines) * 27
        body_height = len(body_lines) * 22
        total = title_height + body_height + 8
        y = y1 + ((y2 - y1 - total) / 2)
        for line in title_lines:
            width, height = text_size(draw, line, BOX_FONT)
            draw.text((x1 + ((x2 - x1 - width) / 2), y), line, font=BOX_FONT, fill=PALETTE["ink"])
            y += height + 6
        y += 2
        for line in body_lines:
            width, height = text_size(draw, line, SMALL_FONT)
            draw.text((x1 + ((x2 - x1 - width) / 2), y), line, font=SMALL_FONT, fill=PALETTE["muted"])
            y += height + 5
    else:
        draw_centered_text(draw, rect, title_lines, BOX_FONT)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color=PALETTE["muted"]) -> None:
    draw.line((start, end), fill=color, width=4)
    sx, sy = start
    ex, ey = end
    dx = ex - sx
    dy = ey - sy
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux = dx / length
    uy = dy / length
    size = 14
    left = (ex - ux * size - uy * size * 0.55, ey - uy * size + ux * size * 0.55)
    right = (ex - ux * size + uy * size * 0.55, ey - uy * size - ux * size * 0.55)
    draw.polygon([end, left, right], fill=color)


def diagram_canvas(title: str, height: int = 1550) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1200, height), PALETTE["surface"])
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((34, 34, 1166, height - 34), radius=34, fill=PALETTE["white"], outline=PALETTE["line"], width=2)
    draw.text((70, 62), title, font=TITLE_FONT, fill=PALETTE["ink"])
    draw.line((70, 116, 1130, 116), fill=PALETTE["line"], width=3)
    return image, draw


def save_diagram(image: Image.Image, name: str) -> Path:
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / name
    image.save(path, "PNG")
    return path


def system_architecture() -> Path:
    image, draw = diagram_canvas("Updated Overall System Architecture", 1500)
    box(draw, (90, 170, 360, 310), "Museum Visitor", "Signs in, browses, scans, listens and navigates", accent=PALETTE["blue"], wrap=18)
    box(draw, (470, 150, 820, 335), "YSMA Flutter Mobile App", "Splash, Auth, Home, Scan, Preview, Details, Map, Profile and Accessibility", accent=PALETTE["gold"], wrap=24)
    box(draw, (900, 170, 1110, 310), "Museum Management", "Uses external Supabase analytics dashboard", accent=PALETTE["purple"], wrap=16)
    arrow(draw, (360, 240), (470, 240))
    arrow(draw, (820, 240), (900, 240))

    modules = [
        ((100, 430, 360, 560), "Camera + TFLite", "50 artwork classes + Unknown; strict confidence, gap, crop and scene-quality gates", PALETTE["red"]),
        ((470, 430, 730, 560), "AR Confirmation", "Full approval mascot appears first, then animated match card and preview", PALETTE["gold"]),
        ((840, 430, 1100, 560), "Artwork Experience", "Preview, detail panel, 43% description area, audio guide and related artworks", PALETTE["green"]),
        ((100, 690, 360, 820), "Museum Map", "Gallery zones, centered pins, route line, destination list and voice route", PALETTE["blue"]),
        ((470, 690, 730, 820), "Accessibility", "Text scale, contrast, theme, audio narration and connected voice navigation", PALETTE["purple"]),
        ((840, 690, 1100, 820), "Profile Activity", "Visitor-facing views, shares, favourites and recent activity only", PALETTE["green"]),
    ]
    for rect, title, body, accent in modules:
        box(draw, rect, title, body, accent=accent, wrap=21)
        arrow(draw, (645, 335), ((rect[0] + rect[2]) // 2, rect[1]))

    box(draw, (105, 1010, 350, 1155), "Shared Preferences", "Local settings, favourites and recent activity", accent=PALETTE["green"], wrap=18)
    box(draw, (465, 990, 735, 1180), "Supabase Backend", "Auth, artworks table, hosted image/audio URLs and admin_activity_events", accent=PALETTE["blue"], wrap=20)
    box(draw, (850, 1010, 1095, 1155), "Dashboard Views", "Overview, top artworks, daily activity and recent activity", accent=PALETTE["purple"], wrap=18)
    arrow(draw, (600, 820), (600, 990))
    arrow(draw, (350, 1082), (465, 1082))
    arrow(draw, (735, 1082), (850, 1082))
    draw.text((115, 1290), "Key update: management analytics is outside the visitor app and reads Supabase dashboard views only.", font=BODY_FONT, fill=PALETTE["ink"])
    return save_diagram(image, "01_updated_system_architecture.png")


def system_flowchart() -> Path:
    image, draw = diagram_canvas("Updated System Flowchart", 1700)
    nodes = [
        ((430, 150, 770, 240), "Start app", "Short splash loads the experience", PALETTE["gold"]),
        ((430, 285, 770, 385), "Authenticate visitor", "Sign in or complete onboarding", PALETTE["blue"]),
        ((430, 430, 770, 530), "Home screen", "Browse, search, scan, map, spotlight or profile", PALETTE["green"]),
        ((95, 620, 385, 745), "Scan artwork", "Camera captures image; TFLite predicts label", PALETTE["red"]),
        ((455, 620, 745, 745), "Use map", "Select gallery/artwork and hear voice route", PALETTE["blue"]),
        ((815, 620, 1105, 745), "Browse details", "Preview, audio guide, description and related works", PALETTE["green"]),
        ((95, 865, 385, 1010), "Reject non-artwork", "Unknown, weak confidence, small gap, low scene quality or unstable match", PALETTE["red"]),
        ((455, 865, 745, 1010), "Confirm real artwork", "Two stable scans match one of the 50 artworks", PALETTE["gold"]),
        ((455, 1130, 745, 1265), "Mascot first", "Approval mascot appears fully with thumbs-up", PALETTE["gold"]),
        ((455, 1370, 745, 1510), "AR-style overlay", "Match card, artwork preview and navigation to detail", PALETTE["purple"]),
        ((815, 1370, 1105, 1510), "Log activity", "Views, shares and favourites sync to Supabase analytics table", PALETTE["blue"]),
    ]
    for rect, title, body, accent in nodes:
        box(draw, rect, title, body, accent=accent, wrap=20)
    arrow(draw, (600, 240), (600, 285))
    arrow(draw, (600, 385), (600, 430))
    arrow(draw, (505, 530), (240, 620))
    arrow(draw, (600, 530), (600, 620))
    arrow(draw, (695, 530), (960, 620))
    arrow(draw, (240, 745), (240, 865))
    arrow(draw, (385, 930), (455, 930))
    arrow(draw, (600, 745), (600, 865))
    arrow(draw, (600, 1010), (600, 1130))
    arrow(draw, (600, 1265), (600, 1370))
    arrow(draw, (745, 1440), (815, 1440))
    arrow(draw, (960, 745), (960, 1370))
    draw.text((110, 1588), "Reject path: if the scene is not one of the 50 artworks, the app keeps scanning instead of forcing a match.", font=SMALL_FONT, fill=PALETTE["muted"])
    return save_diagram(image, "02_updated_system_flowchart.png")


def data_flow_diagram() -> Path:
    image, draw = diagram_canvas("Updated Data Flow Diagram", 1550)
    box(draw, (80, 170, 330, 305), "Visitor", "Credentials, scans, searches, map selections, settings and interactions", accent=PALETTE["blue"], wrap=18)
    box(draw, (455, 145, 745, 330), "Mobile Application", "Flutter UI, recognition gate, AR feedback, audio, map and profile modules", accent=PALETTE["gold"], wrap=21)
    box(draw, (855, 170, 1135, 305), "Museum\nManagement", "Views analytics outside the app", accent=PALETTE["purple"], wrap=18)
    arrow(draw, (330, 237), (455, 237))
    arrow(draw, (745, 237), (855, 237))

    stores = [
        ((70, 500, 350, 665), "Device Camera", "Image bytes for recognition only", PALETTE["red"]),
        ((460, 500, 740, 665), "TFLite Model", "Artwork or Unknown prediction + confidence checks", PALETTE["red"]),
        ((850, 500, 1130, 665), "Supabase", "Auth, artworks, images, audio URLs and analytics events", PALETTE["blue"]),
        ((70, 840, 350, 1005), "Local Storage", "Preferences, favourites and visitor recent activity", PALETTE["green"]),
        ((460, 840, 740, 1005), "Audio/TTS Output", "Artwork audio guide and map voice navigation", PALETTE["purple"]),
        ((850, 840, 1130, 1005), "Admin SQL Views", "Overview, daily activity, top artworks and recent activity", PALETTE["purple"]),
    ]
    for rect, title, body, accent in stores:
        box(draw, rect, title, body, accent=accent, wrap=20)

    arrow(draw, (520, 330), (210, 500))
    arrow(draw, (600, 330), (600, 500))
    arrow(draw, (690, 330), (990, 500))
    arrow(draw, (520, 330), (210, 840))
    arrow(draw, (990, 665), (990, 840))

    labels = [
        (370, 465, "captured frame"),
        (620, 455, "prediction result"),
        (755, 455, "artwork data / event logs"),
        (300, 810, "settings + activity"),
        (875, 790, "structured analytics"),
    ]
    for x, y, text in labels:
        draw.text((x, y), text, font=SMALL_FONT, fill=PALETTE["muted"])
    draw.text((110, 1245), "Privacy update: visitors can generate analytics events, but the dashboard itself is not accessible inside the mobile app.", font=BODY_FONT, fill=PALETTE["ink"])
    return save_diagram(image, "03_updated_data_flow_diagram.png")


def production_workflow() -> Path:
    image, draw = diagram_canvas("Updated Production Workflow", 1700)
    steps = [
        ("Research museum problem", "Study existing labels, navigation limits, accessibility gaps and analytics need", PALETTE["blue"]),
        ("Plan information architecture", "Define visitor screens, scan flow, map flow, profile and external management dashboard", PALETTE["green"]),
        ("Prepare content and assets", "Artwork metadata, images, audio narration, logo, splash assets and training images", PALETTE["gold"]),
        ("Build backend", "Supabase Auth, artworks table, image/audio URLs and analytics event table/views", PALETTE["blue"]),
        ("Create recognition pipeline", "Augment 50 artwork classes, export TFLite model and add unknown/open-set gates", PALETTE["red"]),
        ("Implement Flutter app", "Home, scan, AR feedback, preview, details, audio, map, profile and accessibility", PALETTE["green"]),
        ("Integrate voice and analytics", "Android TTS channel, semantics fallback, local activity and Supabase event logging", PALETTE["purple"]),
        ("Test and refine", "Check false matches, mascot reveal, splash timing, map pins, audio, settings and document flowcharts", PALETTE["gold"]),
    ]
    x1, x2 = 230, 970
    y = 150
    previous_center = None
    for index, (title, body, accent) in enumerate(steps, start=1):
        rect = (x1, y, x2, y + 130)
        box(draw, rect, f"{index}. {title}", body, accent=accent, wrap=32)
        center = ((x1 + x2) // 2, y + 130)
        if previous_center:
            arrow(draw, previous_center, ((x1 + x2) // 2, y))
        previous_center = center
        y += 185
    return save_diagram(image, "04_updated_production_workflow.png")


def replace_text(document: Document, replacements: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text
        if text in replacements:
            paragraph.clear()
            paragraph.add_run(replacements[text])


def clear_table(table) -> None:
    for row in list(table.rows):
        table._tbl.remove(row._tr)


def add_table_row(table, values: list[str], header: bool = False) -> None:
    row = table.add_row()
    for index, value in enumerate(values):
        cell = row.cells[index]
        cell.text = value
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9 if not header else 10)
                run.font.name = "Arial"
                if header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        shading = "13283E" if header else ("F8FAFC" if len(table.rows) % 2 == 0 else "FFFFFF")
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = deepcopy(tc_pr.xpath("./w:shd")[0]) if tc_pr.xpath("./w:shd") else None
        if shd is None:
            from docx.oxml import OxmlElement

            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", shading)


def update_tools_table(document: Document) -> None:
    table = document.tables[0]
    clear_table(table)
    rows = [
        ["Category", "Tools Used in the Production"],
        ["Planning and UI design", "Figma; app flow planning; museum content mapping"],
        ["Mobile framework and language", "Flutter SDK; Dart; Material 3 widgets; CustomPainter and animation controllers"],
        ["Development environment", "Visual Studio Code; Android Studio; Flutter CLI; Android SDK; Gradle"],
        ["Backend and database", "Supabase; PostgreSQL; Supabase Auth; SQL views for external management analytics"],
        ["Artwork recognition", "TensorFlow Lite model; tflite_flutter; image package preprocessing; 50 artwork classes plus Unknown class; training augmentation script"],
        ["Camera and media loading", "camera package; cached_network_image; network image pre-caching"],
        ["Audio and voice navigation", "just_audio; Android TextToSpeech; Flutter MethodChannel; SemanticsService fallback; Audacity for narration preparation"],
        ["State and local storage", "Provider; SharedPreferences for settings, favourites and recent activity"],
        ["Authentication", "Supabase Auth; Google Sign-In dependency"],
        ["AR-style experience", "Flutter animation system; AR-style recognition overlay; mascot approval animation; ar_flutter_plugin dependency for AR expansion"],
        ["Accessibility tools", "Text scaling, high contrast, light/dark mode, audio narration and connected voice navigation"],
        ["Analytics dashboard", "Supabase admin_activity_events table; overview, artwork engagement, daily activity and recent activity SQL views"],
        ["Version control and asset management", "Git; GitHub workflow; Android launcher and splash resources"],
        ["Testing and verification", "Android smartphone testing; Flutter analyzer/build tooling; camera recognition tests; document render QA"],
        ["Documentation", "Microsoft Word; python-docx; Pillow-generated updated flowcharts"],
    ]
    for index, row in enumerate(rows):
        add_table_row(table, row, header=index == 0)


def replace_diagrams(document: Document, diagrams: list[Path]) -> None:
    drawing_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph._p.xpath(".//w:drawing")
    ]
    if len(drawing_paragraphs) != len(diagrams):
        raise RuntimeError(f"Expected {len(diagrams)} diagrams, found {len(drawing_paragraphs)}")

    for paragraph, diagram in zip(drawing_paragraphs, diagrams):
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(diagram), width=Inches(6.25))


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    diagrams = [
        system_architecture(),
        system_flowchart(),
        data_flow_diagram(),
        production_workflow(),
    ]

    document = Document(str(SOURCE))
    replacements = {
        "YSMA Smart Experience serves as a mobile application that merges augmented reality technology with virtual navigation systems to create accessible personalized content while offering fundamental analytic capabilities. The system enables museum visitors who spend time at the museum to use digital support tools which enhance their enjoyment, learning opportunities, and access to museum facilities.":
            "YSMA Smart Experience serves as a mobile application that combines artwork recognition, AR-style feedback, audio interpretation, digital museum navigation, accessibility controls and visitor activity logging. The visitor app is focused on the museum guest experience, while management analytics is handled outside the app through a secured Supabase dashboard structure.",
        "• Providing basic visitor analytics":
            "• Providing basic visitor analytics through an external management dashboard",
        "The proposed system uses a basic client/server model which lets users operate the system through its mobile application client interface. The backend server stores artwork content data, tracks user interactions with the system, and processes analytics. The local database stores artwork metadata and users' interactions with the system in the form of logs.":
            "The proposed system uses a mobile client and cloud backend model. Visitors use the Flutter mobile application to authenticate, browse artworks, scan museum artworks, receive AR-style confirmation, listen to audio guides, use the museum map and adjust accessibility settings. Supabase stores authentication records, artwork metadata, image and audio references and analytics events. Local storage keeps visitor-facing preferences, favourites and recent activity, while management analytics is viewed outside the app through Supabase dashboard views.",
        "1) The Augmented Reality (AR) Module enables users to identify and scan their chosen artworks by using their mobile phone camera to take pictures of the artwork. The system will display extra information through augmented reality after it identifies an artwork. The system will present information through the user's cell phone display, which includes extended descriptions and audio narration, and multimedia content.":
            "1) The Artwork Recognition and AR-style Feedback Module uses the mobile camera and TensorFlow Lite model to identify only the selected museum artworks. The app rejects unknown, weak, unstable or low-quality predictions instead of assigning random objects to artworks. After a valid match, the approval mascot appears fully on screen with a thumbs-up before the animated match card and artwork preview are shown.",
        "2) The Digital Mapping (Virtual Navigation) Module provides users with a digital museum map that shows gallery locations and artwork displays and museum pathways. Users will be able to choose which gallery or piece of art to find on the digital map and will receive visual directions, as well as audio directions through the AR module, to locate the gallery/ piece of art.":
            "2) The Digital Mapping and Voice Navigation Module provides a museum map with gallery zones, centered destination pins, route drawing, destination selection and spoken route instructions. The voice navigation setting in accessibility now connects directly to the same map guidance service.",
        "3) The Accessibility/Adaptive use Module provides multiple ways for users to experience the museum through audio descriptions and enlarged print size and high contrast modes and simplified user interface elements which enable users with different abilities to navigate the museum.":
            "3) The Accessibility/Adaptive Use Module provides text scaling, high contrast mode, light/dark mode, audio narration and voice navigation. These settings support visitors with different visual, reading and navigation needs.",
        "4) The Personalised Recommendation Module uses basic rule-based logic to present users with museum artwork and route recommendations based on their individual preferences and previous interactions with the museum.":
            "4) The Personalised Recommendation and Artwork Experience Module presents recommended artworks, artwork spotlight content, preview pages, detailed descriptions, audio guides and related artworks based on available metadata and visitor interactions.",
        "5) The Analytics Module will store user activities which do not reveal user identity. The system will record activities which include scanned items and navigation requests, and content views for institutional usage.":
            "5) The Analytics Module records visitor activity such as artwork views, shares and favourites. These events are sent to Supabase for museum management and are not displayed inside the visitor app.",
        "The following diagram illustrates that every user action flows through the mobile application prior to their being processed by the backend system and finally being stored or analyzed.":
            "The following updated diagram illustrates the current system architecture, including the visitor app, recognition pipeline, accessibility features, Supabase backend and external management analytics dashboard.",
        "The system flowchart illustrates how a typical visitor interacts with the application.":
            "The system flowchart illustrates the current visitor journey through the implemented application.",
        "The user journey from launching the App through interacting with the system and logging data is shown in this flow.":
            "The user journey from launching the app to scanning, rejecting non-artworks, confirming valid artworks, displaying the mascot, opening AR-style feedback and logging activity is shown in this updated flow.",
        "The Data Flow Diagram shows how data moves through the system.":
            "The Data Flow Diagram shows the updated movement of data between the visitor, mobile application, camera, recognition model, local storage, Supabase backend, audio output and external management dashboard.",
        "This diagram shows that all user actions pass through the mobile app and are processed by the backend system before being stored or analyzed.":
            "This diagram shows that visitor actions pass through the mobile app, while management analytics is stored and viewed outside the visitor application through Supabase.",
        "In this section, we discuss the various tools and technologies utilized in the development of the YSMA Smart Experience Application.":
            "In this section, we discuss the updated tools and technologies used in the development, testing, documentation and deployment preparation of the YSMA Smart Experience application.",
        "The project uses a set of production tools which match the technical requirements and the project scope and the project limitations of the YSMA Smart Experience application. The team selected the tools based on their user-friendliness and their ability to work on mobile devices and their capacity to build augmented reality experiences and their accessibility for users and their ability to create academic prototypes.":
            "The project uses a set of production tools that match the current implemented system. The tools cover mobile development, backend services, artwork recognition, media handling, voice navigation, accessibility, analytics, version control, testing and documentation.",
        "The system development followed a structured production workflow.":
            "The system development followed an updated structured production workflow that reflects the application as implemented.",
        "The workflow proceeds in a natural way from problem identification until system implementation.":
            "The workflow proceeds from museum problem identification through design, content preparation, Supabase setup, recognition model preparation, Flutter implementation, integration, testing and documentation.",
        "The system testing procedures will examine the following components:":
            "The system testing procedures examined the following components:",
        "• The system needs to authenticate the artwork through its recognition capabilities.":
            "• The system needs to authenticate only the 50 selected artworks and reject unknown scenes or random objects.",
        "• The system needs to provide users with clear paths for navigation.":
            "• The system needs to provide users with clear map routes and voice navigation.",
        "• The system needs to evaluate how well it implements its accessibility functions.":
            "• The system needs to evaluate text scaling, contrast mode, audio narration and voice navigation settings.",
        "• The system needs to maintain its operational stability and":
            "• The system needs to maintain operational stability, reduce splash delay and keep interface elements aligned.",
        "• The system needs to record its analytics data correctly.":
            "• The system needs to record analytics events correctly for the external Supabase dashboard.",
    }
    replace_text(document, replacements)
    update_tools_table(document)
    replace_diagrams(document, diagrams)
    document.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
